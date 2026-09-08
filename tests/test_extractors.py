from __future__ import annotations

from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from mcp_sharepoint.extractors import extract_document


def extract(data: bytes, filename: str):
    return extract_document(
        BytesIO(data),
        filename,
        section_character_limit=1_000,
        document_character_limit=10_000,
    )


def test_extract_utf8_text() -> None:
    result = extract("First line\nSecond line".encode(), "notes.txt")

    assert [section.text for section in result.sections] == ["First line\nSecond line"]
    assert result.truncated is False


def test_extract_docx_paragraph_and_table() -> None:
    document = Document()
    document.add_paragraph("Quarterly report")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "42"
    stream = BytesIO()
    document.save(stream)

    result = extract(stream.getvalue(), "report.docx")

    text = "\n".join(section.text for section in result.sections)
    assert "Quarterly report" in text
    assert "Revenue\t42" in text


def test_extract_pptx_by_slide() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Roadmap"
    stream = BytesIO()
    presentation.save(stream)

    result = extract(stream.getvalue(), "roadmap.pptx")

    assert result.sections[0].label == "Slide 1"
    assert "Roadmap" in result.sections[0].text


def test_extract_xlsx_by_sheet() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Forecast"
    worksheet.append(["Month", "Value"])
    worksheet.append(["January", 12])
    stream = BytesIO()
    workbook.save(stream)

    result = extract(stream.getvalue(), "forecast.xlsx")

    assert result.sections[0].label == "Sheet: Forecast"
    assert "January\t12" in result.sections[0].text


def test_long_text_is_split_and_capped() -> None:
    result = extract_document(
        BytesIO(("word " * 1_000).encode()),
        "long.txt",
        section_character_limit=100,
        document_character_limit=250,
    )

    assert all(len(section.text) <= 100 for section in result.sections)
    assert sum(len(section.text) for section in result.sections) <= 250
    assert result.truncated is True
