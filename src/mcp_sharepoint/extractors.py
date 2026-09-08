"""Text extraction for the document formats supported by the MCP server."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from datetime import date, datetime, time
from pathlib import PurePath
from typing import BinaryIO, Iterable

from charset_normalizer import from_bytes
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

SUPPORTED_EXTENSIONS = frozenset(
    {"pdf", "docx", "pptx", "xlsx", "txt", "text", "csv", "md", "log"}
)
_TEXT_EXTENSIONS = frozenset({"txt", "text", "csv", "md", "log"})


class ExtractionError(RuntimeError):
    """Raised when a supported document cannot be read."""


@dataclass(frozen=True, slots=True)
class ExtractedSection:
    index: int
    label: str
    text: str

    def as_dict(self) -> dict[str, int | str]:
        return asdict(self)


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    sections: tuple[ExtractedSection, ...]
    truncated: bool


def extension_for(filename: str) -> str:
    return PurePath(filename).suffix.lower().lstrip(".")


def extract_document(
    stream: BinaryIO,
    filename: str,
    *,
    section_character_limit: int,
    document_character_limit: int,
) -> ExtractionResult:
    """Extract deterministic sections from a supported document."""

    extension = extension_for(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(f"Unsupported file type: .{extension or '(none)'}")

    stream.seek(0)
    try:
        if extension == "pdf":
            labelled_text = _extract_pdf(stream)
        elif extension == "docx":
            labelled_text = _extract_docx(stream)
        elif extension == "pptx":
            labelled_text = _extract_pptx(stream)
        elif extension == "xlsx":
            labelled_text = _extract_xlsx(stream)
        else:
            labelled_text = _extract_plain_text(stream)

        return _build_sections(
            labelled_text,
            section_character_limit=section_character_limit,
            document_character_limit=document_character_limit,
        )
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            f"Could not extract text from {filename!r}: {exc}"
        ) from exc


def _extract_pdf(stream: BinaryIO) -> Iterable[tuple[str, str]]:
    reader = PdfReader(stream, strict=False)
    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise ExtractionError("The PDF is encrypted")
        except Exception as exc:
            raise ExtractionError("The PDF is encrypted") from exc

    return (
        (f"Page {page_number}", page.extract_text() or "")
        for page_number, page in enumerate(reader.pages, start=1)
    )


def _extract_docx(stream: BinaryIO) -> Iterable[tuple[str, str]]:
    document = Document(stream)
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text)

    for table_number, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            values = [_clean_cell(cell.text) for cell in row.cells]
            if any(values):
                rows.append("\t".join(values))
        if rows:
            blocks.append(f"Table {table_number}\n" + "\n".join(rows))

    return (("Document", "\n\n".join(blocks)),)


def _extract_pptx(stream: BinaryIO) -> Iterable[tuple[str, str]]:
    presentation = Presentation(stream)
    slides: list[tuple[str, str]] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if isinstance(text, str) and text.strip():
                blocks.append(text)
            if getattr(shape, "has_table", False):
                rows: list[str] = []
                for row in shape.table.rows:
                    values = [_clean_cell(cell.text) for cell in row.cells]
                    if any(values):
                        rows.append("\t".join(values))
                if rows:
                    blocks.append("\n".join(rows))
        slides.append((f"Slide {slide_number}", "\n\n".join(blocks)))

    return slides


def _extract_xlsx(stream: BinaryIO) -> Iterable[tuple[str, str]]:
    workbook = load_workbook(stream, read_only=True, data_only=True)
    sheets: list[tuple[str, str]] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                values = [_spreadsheet_value(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append("\t".join(values))
            sheets.append((f"Sheet: {worksheet.title}", "\n".join(rows)))
    finally:
        workbook.close()
    return sheets


def _extract_plain_text(stream: BinaryIO) -> Iterable[tuple[str, str]]:
    data = stream.read()
    match = from_bytes(data).best()
    text = str(match) if match is not None else data.decode("utf-8", errors="replace")
    return (("Text", text),)


def _build_sections(
    labelled_text: Iterable[tuple[str, str]],
    *,
    section_character_limit: int,
    document_character_limit: int,
) -> ExtractionResult:
    sections: list[ExtractedSection] = []
    characters_used = 0
    truncated = False

    for label, raw_text in labelled_text:
        text = _normalize_text(raw_text)
        if not text:
            continue
        remaining = document_character_limit - characters_used
        if remaining <= 0:
            truncated = True
            break
        if len(text) > remaining:
            text = text[:remaining]
            truncated = True

        parts = _split_text(text, section_character_limit)
        for part_number, part in enumerate(parts, start=1):
            part_label = label if len(parts) == 1 else f"{label} (part {part_number})"
            sections.append(
                ExtractedSection(index=len(sections), label=part_label, text=part)
            )
            characters_used += len(part)

        if truncated:
            break

    return ExtractionResult(sections=tuple(sections), truncated=truncated)


def _split_text(text: str, limit: int) -> list[str]:
    if len(text) <= limit:
        return [text]

    parts: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + limit, len(text))
        if end < len(text):
            preferred_break = text.rfind("\n", start + limit // 2, end)
            if preferred_break < 0:
                preferred_break = text.rfind(" ", start + limit // 2, end)
            if preferred_break >= 0:
                end = preferred_break
        part = text[start:end].strip()
        if part:
            parts.append(part)
        start = end
        while start < len(text) and text[start].isspace():
            start += 1
    return parts


def _normalize_text(value: str) -> str:
    text = value.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _clean_cell(value: str) -> str:
    return " ".join(value.replace("\r", " ").replace("\n", " ").split())


def _spreadsheet_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return _clean_cell(str(value))
